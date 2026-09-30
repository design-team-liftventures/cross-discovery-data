"""Source adapters.

Each adapter reads the raw exports of one source for one month and hands rows to
the Run: source-table rows, evidence items, rejects, identity links, a masked
mirror copy of every file, and a reconciliation line per file:

    rows_read == rows_loaded + rows_rejected   (or the run fails)

Row numbers are 1-based positions of the data record in the file, after the
header row(s) — not physical line numbers (quoted fields can span lines).
"""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

from common import (dumps, initiative_of, is_blank, parse_dt, read_csv,
                    short_hash, slug, to_float)

MASK = "[name]"


def _pad(row: list[str], n: int) -> list[str]:
    return row + [""] * (n - len(row)) if len(row) < n else row


def _iso(dt):
    return dt.isoformat(sep=" ") if dt else None


# --------------------------------------------------------------------------
# Chargebee (Brightback cancel flow)
# --------------------------------------------------------------------------
CB_NAME_COLS = ["ownerFirstName", "ownerLastName", "ownerFullName",
                "Owner First Name", "Owner Last Name", "Owner Full Name"]
CB_EMAIL_COLS = ["email", "Owner Email"]
CB_HASH_COLS = ["Billing ID", "Subscription ID"]
CB_TEXT_COLS = ["reasonOther", "competitor", "competitionOther", "details",
                "drilldown_modal_feedback", "free_text_modal_feedback"]


def chargebee_names(run, month, files):
    for f in files:
        rows = read_csv(f)
        if not rows:
            continue
        idx = {c: i for i, c in enumerate(rows[0])}
        for row in rows[1:]:
            row = _pad(row, len(rows[0]))
            g = lambda c: row[idx[c]] if c in idx else ""
            run.anon.add_name(g("ownerFirstName"), g("ownerLastName"))
            run.anon.add_name(g("Owner First Name"))
            run.anon.add_name(g("ownerFullName"))


def chargebee(run, month, files):
    anon = run.anon
    for f in files:
        rows = read_csv(f)
        if not rows:
            run.recon("chargebee", month, f, 0, 0, 0)
            continue
        header, data = rows[0], rows[1:]
        run.schema("chargebee", f, header)
        idx = {c: i for i, c in enumerate(header)}
        mirror = [header]
        loaded = rejected = 0
        for n, row in enumerate(data, 1):
            row = _pad(row, len(header))
            g = lambda c: row[idx[c]] if c in idx else ""
            email = g("email") or g("Owner Email")
            uk = anon.user_key(email)
            names = tuple(g(c) for c in CB_NAME_COLS) + tuple(g("ownerFullName").split())
            masked = list(row)
            for c in CB_EMAIL_COLS:
                if c in idx and masked[idx[c]]:
                    masked[idx[c]] = uk or "[email]"
            for c in CB_NAME_COLS:
                if c in idx and masked[idx[c]]:
                    masked[idx[c]] = MASK
            for c in CB_HASH_COLS:
                if c in idx:
                    masked[idx[c]] = anon.hash_id(g(c)) or ""
            for c in CB_TEXT_COLS:
                if c in idx:
                    masked[idx[c]] = anon.scrub(g(c), names)
            mirror.append(masked)

            sid = g("exit_session_id") or f"cb_{short_hash(str(f), str(n))}"
            if is_blank(row):
                run.reject("chargebee", month, f, n, sid, "empty_row"); rejected += 1; continue
            if run.seen("chargebee", sid):
                run.reject("chargebee", month, f, n, sid, "duplicate"); rejected += 1; continue

            # internal_id is the SuperSummary user id (= Amplitude user_id) only when numeric;
            # otherwise it is a Chargebee customer id, kept hashed as billing_customer_ref.
            raw_id = (g("internal_id") or g("id")).strip()
            uid = raw_id if raw_id.isdigit() else None
            billing_ref = anon.hash_id(raw_id, "cb_") if raw_id and not raw_id.isdigit() else None
            cancel = parse_dt(g("Cancel date")) or parse_dt(g("startTime"))
            if not cancel:
                run.alert(f"chargebee: unparsed cancel date in {f.name} row {n}")
            m = lambda c: masked[idx[c]] if c in idx else ""
            extra = {c: masked[i] for c, i in idx.items()
                     if c not in CB_NAME_COLS + CB_EMAIL_COLS + ["id", "internal_id"] and masked[i] != ""}
            rec = {
                "session_id": sid,
                "period": month,
                "cancel_at": _iso(cancel),
                "user_key": uk,
                "user_id": uid,
                "billing_customer_ref": billing_ref,
                "audience": (g("Customer Audience") or "").strip().lower().replace(" ", "_") or None,
                "audiences_names": g("Audiences Names") or None,
                "reason": g("reason") or None,
                "reason_category": g("bbkCategory") or None,
                "reason_other": m("reasonOther") or None,
                "details": m("details") or None,
                "competitor": m("competitor") or None,
                "return_likelihood": to_float(g("returnLikelihood")),
                "accepted_offers": g("acceptedOffers") or None,
                "plan_id": g("Plan ID") or None,
                "plan_interval": g("Plan Interval") or None,
                "billing_interval": g("billing_interval") or None,
                "mrr_usd": to_float(g("mrr")),
                "billing_price_usd": to_float(g("billing_price")),
                "age_days": to_float(g("age")),
                "grade_level": g("Customer Grade Level") or None,
                "page_attributed": g("Page attributed") or None,
                "coupon_ids": g("Coupon IDs") or None,
                "winbacks": to_float(g("winbacks")),
                "validation_status": g("validation_status") or None,
                "subscription_status": g("Subscription Status") or None,
                "first_purchase_at": _iso(parse_dt(g("First Purchase Date"))),
                "last_active_at": _iso(parse_dt(g("lastActiveDate"))),
                "raw": dumps(extra),
                "source_file": run.rel(f),
                "row_number": n,
            }
            run.add("cancellations", rec); loaded += 1
            run.identity(uk, uid, "chargebee")

            parts = [p for p in (rec["reason_other"], rec["details"]) if p]
            if rec["competitor"]:
                parts.append(f"Competitor: {rec['competitor']}")
            run.add_evidence(
                source="chargebee", record_id=sid, period=month, occurred_at=rec["cancel_at"],
                user_key=uk, user_id=uid, instrument="Chargebee cancel flow",
                question="Cancel reason", text=" | ".join(parts) or None,
                category=rec["reason"], score=rec["return_likelihood"], score_type="return_likelihood_0_10",
                audience=rec["audience"], plan=rec["plan_id"], mapping=("chargebee_reason", rec["reason"]),
                segment=("BCM" if rec["audience"] == "book_club_member" else "Non-BCM") if rec["audience"] else None,
                source_file=rec["source_file"], row_number=n)
        run.recon("chargebee", month, f, len(data), loaded, rejected)
        run.write_mirror(month, f, mirror)


# --------------------------------------------------------------------------
# Sprig (NPS, PMF, exit intent, product surveys)
# --------------------------------------------------------------------------
Q_RE = re.compile(r"^Q(\d+)_(.+)$")
SPRIG_PII_ATTRS = {"email", "firstName", "lastName"}


def _sprig_family(path: Path, name: str) -> str:
    p = str(path).lower()
    n = (name or "").lower()
    if "/nps/" in p or "nps" in n:
        return "NPS"
    if "/pmf/" in p or n.startswith("pmf"):
        return "PMF"
    ini = initiative_of(name, path.name)
    return ini or "Other"


def sprig_names(run, month, files):
    for f in files:
        rows = read_csv(f)
        if len(rows) < 2:
            continue
        header, labels = rows[0], rows[1]
        fi = [i for i, c in enumerate(header) if c.startswith("Attribute_") and i < len(labels) and labels[i] == "firstName"]
        li = [i for i, c in enumerate(header) if c.startswith("Attribute_") and i < len(labels) and labels[i] == "lastName"]
        for row in rows[2:]:
            row = _pad(row, len(header))
            if fi and li:
                run.anon.add_name(row[fi[0]], row[li[0]])


def sprig(run, month, files):
    anon = run.anon
    for f in files:
        rows = read_csv(f)
        if len(rows) < 2:
            run.recon("sprig", month, f, 0, 0, 0)
            continue
        header, labels, data = rows[0], _pad(rows[1], len(rows[0])), rows[2:]
        run.schema("sprig", f, header)
        idx = {c: i for i, c in enumerate(header)}
        attr_cols = {i: labels[i] for i, c in enumerate(header) if c.startswith("Attribute_")}
        # question structure
        qs: dict[int, dict] = {}
        for i, c in enumerate(header):
            mq = Q_RE.match(c)
            if not mq:
                continue
            q = qs.setdefault(int(mq.group(1)), {"themes": [], "choices": [], "other": []})
            kind = mq.group(2)
            if kind == "Question_Text":
                q["text"] = i
            elif kind == "Response":
                q["response"] = i
            elif kind == "Value":
                q["value"] = i
            elif kind == "Response_Other":
                q["resp_other"] = i
            elif kind == "Response_NPSCategory":
                q["nps"] = i
            elif kind == "Theme":
                q["theme"] = i
            elif re.fullmatch(r"Theme_\d+", kind):
                q["themes"].append(i)
            elif re.fullmatch(r"Choice_\d+", kind):
                q["choices"].append(i)
            else:
                q["other"].append(i)
        text_cols = {q[k] for q in qs.values() for k in ("response", "resp_other") if k in q}

        mirror = [header, labels]
        loaded = rejected = 0
        for n, row in enumerate(data, 1):
            row = _pad(row, len(header))
            attrs = {lab: row[i] for i, lab in attr_cols.items() if row[i] != ""}
            uk = anon.user_key(attrs.get("email", ""))
            names = (attrs.get("firstName", ""), attrs.get("lastName", ""))
            masked = list(row)
            for i, lab in attr_cols.items():
                if lab == "email" and masked[i]:
                    masked[i] = uk or "[email]"
                elif lab in ("firstName", "lastName") and masked[i]:
                    masked[i] = MASK
            for i in text_cols:
                masked[i] = anon.scrub(row[i], names)
            mirror.append(masked)

            g = lambda c: row[idx[c]] if c in idx else ""
            rid = g("responseGroupUid") or f"sp_{short_hash(str(f), str(n))}"
            if is_blank(row):
                run.reject("sprig", month, f, n, rid, "empty_row"); rejected += 1; continue
            if run.seen("sprig", rid):
                run.reject("sprig", month, f, n, rid, "duplicate"); rejected += 1; continue
            if run.cfg["sprig"].get("exclude_staff", True) and attrs.get("isStaff", "").strip().lower() == "true":
                run.reject("sprig", month, f, n, rid, "staff_response"); rejected += 1; continue

            survey_name = g("surveyName")
            family = _sprig_family(f, survey_name)
            fname = f.name.lower()
            audience = (attrs.get("audienceCategory") or attrs.get("audience") or "").strip().lower().replace(" ", "_") or None
            segment = "Non-BCM" if "non-bcm" in fname else ("BCM" if "bcm" in fname else None)
            if segment is None and audience:
                segment = "BCM" if audience == "book_club_member" else "Non-BCM"
            channel = "app" if (attrs.get("appVersion") or attrs.get("sdkVersion")) else "web"
            uid = g("userId").strip() or None
            created = parse_dt(g("createdAt"))
            if g("createdAt") and not created:
                run.alert(f"sprig: unparsed createdAt in {f.name} row {n}")
            safe_attrs = {k: v for k, v in attrs.items() if k not in SPRIG_PII_ATTRS}
            rec = {
                "response_id": rid,
                "period": month,
                "survey_id": g("surveyId") or None,
                "survey_name": survey_name or None,
                "survey_family": family,
                "initiative": initiative_of(survey_name, f.name),
                "created_at": _iso(created),
                "completed_at": _iso(parse_dt(g("completedAt"))),
                "user_key": uk,
                "user_id": uid,
                "visitor_id": g("visitorId") or None,
                "segment": segment,
                "audience": audience,
                "subscriber_status": attrs.get("subscriberStatus") or attrs.get("status") or None,
                "plan": attrs.get("planId") or attrs.get("planType") or None,
                "tenure": attrs.get("tenure") or None,
                "grade_level": attrs.get("gradeLevel") or None,
                "channel": channel,
                "os": g("os") or None,
                "browser": g("browser") or None,
                "device_type": attrs.get("deviceType") or None,
                "href": g("href") or None,
                "triggering_event": g("triggeringEvent") or None,
                "attributes": dumps(safe_attrs),
                "source_file": run.rel(f),
                "row_number": n,
            }
            run.add("sprig_responses", rec); loaded += 1
            run.identity(uk, uid, "sprig")

            for qn in sorted(qs):
                q = qs[qn]
                resp = masked[q["response"]] if "response" in q else ""
                other = masked[q["resp_other"]] if "resp_other" in q else ""
                qtext = (row[q["text"]] if "text" in q else "") or (labels[q["text"]] if "text" in q else "")
                qtext = " ".join(qtext.split())
                themes = [labels[i] for i in q["themes"] if row[i] not in ("", "0", "false", "False")]
                choices = []
                for i in q["choices"]:
                    v = row[i]
                    if v in ("", "0", "false", "False"):
                        continue
                    choices.append(labels[i] if v in ("1", "true", "True") and labels[i] else v)
                primary = row[q["theme"]] if "theme" in q else ""
                value = to_float(row[q["value"]]) if "value" in q else None
                nps_cat = row[q["nps"]] if "nps" in q else ""
                if not (resp or other or themes or choices):
                    continue
                skipped = resp.strip().lower() == "[skipped by user]"
                if skipped:
                    resp = ""
                score, score_type, category = None, None, None
                ql = qtext.lower()
                if "nps" in q:
                    score, score_type, category = to_float(resp), "nps_0_10", nps_cat or None
                elif family == "PMF" and ("no longer use" in ql or "feel if" in ql):
                    # Sean Ellis question on a 1-5 scale; 5 = very disappointed (matches the PMF report)
                    score = to_float(resp)
                    score_type = "pmf_1_5"
                    if score is not None:
                        category = {5: "very_disappointed", 4: "somewhat_disappointed"}.get(int(score), "not_disappointed")
                    elif resp:
                        category = "very_disappointed" if "very" in resp.lower() else resp
                elif value is not None:
                    score, score_type = value, "choice_value"
                elif to_float(resp) is not None and not (themes or choices):
                    score, score_type = to_float(resp), "scale"
                ans = {
                    "response_id": rid, "question_no": qn, "question": qtext or None, "skipped": skipped,
                    "response": resp or None, "response_other": other or None,
                    "value": value, "nps_category": nps_cat or None,
                    "primary_theme": primary or None, "themes": themes, "choices": choices,
                }
                run.add("sprig_answers", ans)
                if skipped and not (other or themes or choices):
                    continue
                text = resp or ""
                if other:
                    text = f"{text} | Other: {other}" if text else other
                if choices and not text:
                    text = "; ".join(choices)
                run.add_evidence(
                    source="sprig", record_id=rid, sub_id=str(qn), period=month, occurred_at=rec["created_at"],
                    user_key=uk, user_id=uid, instrument=survey_name or f.stem, question=qtext or None,
                    text=text or None, category=category, score=score, score_type=score_type,
                    source_theme="; ".join([primary] + themes if primary else themes),
                    segment=segment, audience=rec["audience"], channel=channel, plan=rec["plan"],
                    initiative=rec["initiative"], source_file=rec["source_file"], row_number=n,
                    choice=bool(q["choices"]) or ("value" in q and "nps" not in q))
        run.recon("sprig", month, f, len(data), loaded, rejected)
        run.write_mirror(month, f, mirror)


# --------------------------------------------------------------------------
# Freshdesk
# --------------------------------------------------------------------------
CUSTOMER_FREEMAIL = None  # all non-system senders count as customers


def _sender_type(addr: str, cfg: dict) -> str:
    a = (addr or "").strip().lower()
    dom = a.split("@")[-1]
    fd = cfg["freshdesk"]
    for d in fd["exclude_internal_domains"]:
        if dom == d or dom.endswith("." + d):
            return "internal"
    for p in fd["exclude_automated_patterns"]:
        if p in a:
            return "automated"
    for t, doms in fd["sender_types"].items():
        for d in doms:
            if dom == d or dom.endswith("." + d):
                return t
    return "customer"


def freshdesk(run, month, files):
    anon = run.anon
    for f in files:
        rows = read_csv(f)
        if not rows:
            run.recon("freshdesk", month, f, 0, 0, 0)
            continue
        header, data = rows[0], rows[1:]
        run.schema("freshdesk", f, header)
        idx = {c: i for i, c in enumerate(header)}
        need = ["Date", "Ticket ID", "User", "Segmentation", "Feedback"]
        missing = [c for c in need if c not in idx]
        if missing:
            run.alert(f"freshdesk: {f.name} missing columns {missing}")
        mirror = [header]
        loaded = rejected = 0
        for n, row in enumerate(data, 1):
            row = _pad(row, len(header))
            g = lambda c: row[idx[c]] if c in idx else ""
            addr = g("User").strip()
            stype = _sender_type(addr, run.cfg)
            uk = anon.user_key(addr) if stype == "customer" else None
            local = addr.split("@")[0]
            own = tuple(t for t in re.split(r"[._\-+0-9]+", local) if len(t) >= 3)
            subject = anon.scrub(g("Segmentation"), own, aggressive=True)
            subject = re.sub(r"^((?:Re:\s*)?(?:Cancel(?:lation)?|Refund)[^-–—]*[-–—]\s*)([A-Z][a-z]+(?:\s+[A-Z][a-zA-Z'\-]+)?)\s*$",
                             lambda m: m.group(1) + MASK, subject)
            body = anon.scrub(g("Feedback"), own, aggressive=True)
            masked = list(row)
            if "User" in idx:
                masked[idx["User"]] = uk or ("@" + addr.split("@")[-1].lower() if "@" in addr else "")
            if "Segmentation" in idx:
                masked[idx["Segmentation"]] = subject
            if "Feedback" in idx:
                masked[idx["Feedback"]] = body
            mirror.append(masked)

            tid = g("Ticket ID").strip() or f"fd_{short_hash(str(f), str(n))}"
            if is_blank(row):
                run.reject("freshdesk", month, f, n, tid, "empty_row"); rejected += 1; continue
            if run.seen("freshdesk", tid):
                run.reject("freshdesk", month, f, n, tid, "duplicate"); rejected += 1; continue
            if stype in ("internal", "automated"):
                run.reject("freshdesk", month, f, n, tid, f"sender_{stype}"); rejected += 1; continue
            created = parse_dt(g("Date"))
            if g("Date") and not created:
                run.alert(f"freshdesk: unparsed Date in {f.name} row {n}")
            rec = {
                "ticket_id": tid, "period": month, "created_at": _iso(created),
                "sender_type": stype, "sender_domain": addr.split("@")[-1].lower() if "@" in addr else None,
                "user_key": uk, "subject": subject or None, "body": body or None,
                "body_chars": len(body or ""), "source_file": run.rel(f), "row_number": n,
            }
            run.add("tickets", rec); loaded += 1
            text = "\n".join(x for x in (subject, body) if x)
            run.add_evidence(
                source="freshdesk", record_id=tid, period=month, occurred_at=rec["created_at"],
                user_key=uk, user_id=None, instrument="Freshdesk ticket", question=None,
                text=text or None, category=stype, mapping=("freshdesk_sender_type", stype),
                source_file=rec["source_file"], row_number=n)
        run.recon("freshdesk", month, f, len(data), loaded, rejected)
        run.write_mirror(month, f, mirror)


# --------------------------------------------------------------------------
# Money-Back Guarantee form
# --------------------------------------------------------------------------
MBG_MAP = [
    (re.compile(r"^\s*$"), "status"),
    (re.compile(r"notes", re.I), "notes_sent_email"),
    (re.compile(r"^amount$", re.I), "amount_usd"),
    (re.compile(r"eligible", re.I), "refund_outcome"),
    (re.compile(r"timestamp", re.I), "requested_at"),
    (re.compile(r"email", re.I), "email"),
    (re.compile(r"made you stay", re.I), "would_have_stayed_if"),
    (re.compile(r"why do you want to cancel", re.I), "why_cancel"),
    (re.compile(r"not pleased", re.I), "not_pleased_because"),
    (re.compile(r"title and author", re.I), "title_and_author"),
    (re.compile(r"i think the content", re.I), "content_opinion"),
    (re.compile(r"expected to find", re.I), "expected_to_find"),
    (re.compile(r"technical issue", re.I), "technical_issue"),
    (re.compile(r"don.?t need it anymore", re.I), "dont_need_because"),
    (re.compile(r"dropdown", re.I), "dropdown"),
]
MBG_TEXT = ["would_have_stayed_if", "why_cancel", "not_pleased_because", "title_and_author",
            "content_opinion", "expected_to_find", "technical_issue", "dont_need_because", "dropdown"]


def mbg(run, month, files):
    anon = run.anon
    for f in files:
        rows = read_csv(f)
        if not rows:
            run.recon("mbg", month, f, 0, 0, 0)
            continue
        header, data = rows[0], rows[1:]
        run.schema("mbg", f, header)
        colmap = {}
        for i, c in enumerate(header):
            for rx, name in MBG_MAP:
                if rx.search(c) and name not in colmap:
                    colmap[name] = i
                    break
            else:
                run.alert(f"mbg: unmapped column '{c}' in {f.name} (kept in mirror only)")
        mirror = [header]
        loaded = rejected = 0
        for n, row in enumerate(data, 1):
            row = _pad(row, len(header))
            g = lambda k: row[colmap[k]] if k in colmap else ""
            uk = anon.user_key(g("email"))
            masked = list(row)
            if "email" in colmap and masked[colmap["email"]]:
                masked[colmap["email"]] = uk or "[email]"
            for k in MBG_TEXT + ["notes_sent_email"]:
                if k in colmap:
                    masked[colmap[k]] = anon.scrub(row[colmap[k]])
            mirror.append(masked)
            mm = lambda k: masked[colmap[k]] if k in colmap else ""
            req = parse_dt(g("requested_at"))
            rid = "mbg_" + short_hash((uk or ""), g("requested_at"), g("why_cancel"))
            if is_blank(row):
                run.reject("mbg", month, f, n, rid, "empty_row"); rejected += 1; continue
            if run.seen("mbg", rid):
                run.reject("mbg", month, f, n, rid, "duplicate"); rejected += 1; continue
            if g("requested_at") and not req:
                run.alert(f"mbg: unparsed Timestamp in {f.name} row {n}")
            rec = {"request_id": rid, "period": month, "requested_at": _iso(req), "user_key": uk,
                   "status": g("status").strip() or None, "refund_outcome": g("refund_outcome") or None,
                   "amount_usd": to_float(g("amount_usd")), "notes_sent_email": mm("notes_sent_email") or None}
            for k in MBG_TEXT:
                rec[k] = mm(k) or None
            rec["source_file"] = run.rel(f); rec["row_number"] = n
            run.add("refunds", rec); loaded += 1
            labels = {"why_cancel": "Why cancel", "would_have_stayed_if": "Would have stayed if",
                      "not_pleased_because": "Not pleased because", "content_opinion": "Content",
                      "expected_to_find": "Expected to find", "technical_issue": "Technical issue",
                      "dont_need_because": "Don't need it because", "title_and_author": "Title"}
            text = " | ".join(f"{labels[k]}: {rec[k]}" for k in labels if rec.get(k))
            run.add_evidence(
                source="mbg", record_id=rid, period=month, occurred_at=rec["requested_at"], user_key=uk,
                user_id=None, instrument="Money-Back Guarantee form", question="Why do you want to cancel?",
                text=text or None, category=rec["why_cancel"], score=rec["amount_usd"], score_type="refund_amount_usd",
                source_file=rec["source_file"], row_number=n)
        run.recon("mbg", month, f, len(data), loaded, rejected)
        run.write_mirror(month, f, mirror)


# --------------------------------------------------------------------------
# Lyssna (panel tests)
# --------------------------------------------------------------------------
LYSSNA_META = ["Response ID", "Date", "Duration (ms)", "Age", "Country", "Gender", "Education level",
               "Annual household income", "Employment status", "Department", "Job", "Daily hours online",
               "Technical proficiency", "General hobbies and interests", "Device Type", "Device Platform",
               "Device Width", "Device Height"]
LY_COL_RE = re.compile(r'^(?P<section>\d+(?:\.\d+)?\.?\s*\([^)]*\)|\([^)]*\))?\s*(?:\((?P<qlabel>[^:()]+):\s*"(?P<q>.*?)"\))?\s*(?P<field>.*)$', re.S)


def _lyssna_study_name(f: Path) -> str:
    return re.sub(r"\s*-\s*results.*$|-results.*$", "", f.stem).strip()


def lyssna(run, month, files):
    anon = run.anon
    for f in files:
        rows = read_csv(f)
        if not rows:
            run.recon("lyssna", month, f, 0, 0, 0)
            continue
        header, data = rows[0], rows[1:]
        run.schema("lyssna", f, header)
        name = _lyssna_study_name(f)
        study_id = "lys_" + short_hash(month, name)
        meta_idx = {c: i for i, c in enumerate(header) if c in LYSSNA_META}
        q_idx = [i for i, c in enumerate(header) if c not in LYSSNA_META]
        parsed = {}
        for i in q_idx:
            m = LY_COL_RE.match(header[i])
            parsed[i] = {
                "section": (m.group("section") or "").strip() if m else "",
                "question": (m.group("q") or "").strip() if m else "",
                "field": (m.group("field") or "").strip() if m else header[i],
            }
        mirror = [header]
        loaded = rejected = 0
        for n, row in enumerate(data, 1):
            row = _pad(row, len(header))
            masked = list(row)
            for i in q_idx:
                masked[i] = anon.scrub(row[i])
            mirror.append(masked)
            rid = row[meta_idx["Response ID"]] if "Response ID" in meta_idx else f"ly_{short_hash(str(f), str(n))}"
            pid = f"{study_id}:{rid}"
            if is_blank(row):
                run.reject("lyssna", month, f, n, pid, "empty_row"); rejected += 1; continue
            if run.seen("lyssna", pid):
                run.reject("lyssna", month, f, n, pid, "duplicate"); rejected += 1; continue
            gm = lambda c: row[meta_idx[c]] if c in meta_idx else ""
            date = parse_dt(gm("Date"))
            rec = {"study_id": study_id, "response_id": rid, "period": month, "responded_at": _iso(date),
                   "duration_ms": to_float(gm("Duration (ms)")), "country": gm("Country") or None,
                   "age": gm("Age") or None, "gender": gm("Gender") or None,
                   "device_type": gm("Device Type") or None, "device_platform": gm("Device Platform") or None,
                   "demographics": dumps({c: gm(c) for c in meta_idx if gm(c)}),
                   "source_file": run.rel(f), "row_number": n}
            run.add("lyssna_participants", rec); loaded += 1
            for i in q_idx:
                v = masked[i]
                if not v:
                    continue
                p = parsed[i]
                run.add("lyssna_answers", {"study_id": study_id, "response_id": rid, "column_index": i,
                                           "column": header[i], "section": p["section"] or None,
                                           "question": p["question"] or None, "field": p["field"] or None,
                                           "answer": v})
                field = (p["field"] or "").lower()
                textual = (len(v) >= 15 and " " in v and not v.startswith("http")
                           and to_float(v) is None and ("answer" in field or field == "" or "response" in field))
                if textual:
                    run.add_evidence(
                        source="lyssna", record_id=pid, sub_id=str(i), period=month, occurred_at=rec["responded_at"],
                        user_key=None, user_id=None, instrument=name, question=p["question"] or header[i][:200],
                        text=v, population="panel", initiative=initiative_of(name),
                        channel=(gm("Device Type") or None), source_file=rec["source_file"], row_number=n)
        run.add_study({"study_id": study_id, "kind": "lyssna_test", "name": name, "initiative": initiative_of(name),
                       "period": month, "n": loaded, "source_file": run.rel(f)})
        run.recon("lyssna", month, f, len(data), loaded, rejected)
        run.write_mirror(month, f, mirror)


# --------------------------------------------------------------------------
# Interviews tracker (Notion export)
# --------------------------------------------------------------------------
IV_DROP = {"Interview recording link", "Transcript link", "Snapshot link", "Metabase Info"}


def interviews_names(run, month, files):
    for f in files:
        rows = read_csv(f)
        if not rows:
            continue
        idx = {c: i for i, c in enumerate(rows[0])}
        for row in rows[1:]:
            row = _pad(row, len(rows[0]))
            if "Users" in idx:
                run.anon.add_name(row[idx["Users"]])


def interviews(run, month, files):
    anon = run.anon
    for f in files:
        rows = read_csv(f)
        if not rows:
            run.recon("interviews", month, f, 0, 0, 0)
            continue
        header, data = rows[0], rows[1:]
        run.schema("interviews", f, header)
        idx = {c: i for i, c in enumerate(header)}
        mirror = [header]
        loaded = rejected = 0
        for n, row in enumerate(data, 1):
            row = _pad(row, len(header))
            g = lambda c: row[idx[c]] if c in idx else ""
            uk = anon.user_key(g("Email"))
            own = tuple(g("Users").split()) + (g("Users"),)
            masked = list(row)
            for c, i in idx.items():
                if c == "Users" and masked[i]:
                    masked[i] = MASK
                elif c == "Email" and masked[i]:
                    masked[i] = uk or "[email]"
                elif c in IV_DROP and masked[i]:
                    masked[i] = "[removed]"
                else:
                    masked[i] = anon.scrub(row[i], own, aggressive=True)
            mirror.append(masked)
            m = lambda c: masked[idx[c]] if c in idx else ""
            iid = "iv_" + short_hash(uk or "", g("Interview date"), g("Booking date"), str(n))
            if is_blank(row):
                run.reject("interviews", month, f, n, iid, "empty_row"); rejected += 1; continue
            idate = parse_dt(g("Interview date"))
            rec = {"interview_id": iid, "period": month, "user_key": uk, "segment": g("Segment") or None,
                   "booking_date": _iso(parse_dt(g("Booking date"))), "interview_date": _iso(idate),
                   "lead_interviewer": g("Lead Interviewer") or None, "outcome": g("Outcome") or None,
                   "plan": g("Plan") or None, "snapshot": m("Snapshot (for Zapier)") or None,
                   "transcript": m("Transcript (for Zapier)") or None,
                   "notes": m("Additional Notes") or None,
                   "raw": dumps({c: masked[i] for c, i in idx.items() if masked[i] and c not in
                                 ("Snapshot (for Zapier)", "Transcript (for Zapier)")}),
                   "source_file": run.rel(f), "row_number": n}
            run.add("interviews", rec); loaded += 1
            if rec["snapshot"]:
                run.add_evidence(
                    source="interviews", record_id=iid, period=month, occurred_at=rec["interview_date"],
                    user_key=uk, user_id=None, instrument="User interview", question="Interview snapshot",
                    text=rec["snapshot"], segment=rec["segment"], plan=rec["plan"],
                    initiative=initiative_of(rec["snapshot"][:300]), source_file=rec["source_file"], row_number=n)
        run.add_study({"study_id": "ivt_" + short_hash(month, f.name), "kind": "interviews", "name": f"Interviews tracker ({f.name})",
                       "initiative": None, "period": month, "n": loaded, "source_file": run.rel(f)})
        run.recon("interviews", month, f, len(data), loaded, rejected)
        run.write_mirror(month, f, mirror)


# --------------------------------------------------------------------------
# Documents (reports, syntheses): docx / pdf / md / txt
# --------------------------------------------------------------------------
def _extract_text(f: Path) -> str:
    ext = f.suffix.lower()
    if ext in (".md", ".txt"):
        return f.read_text(encoding="utf-8", errors="replace")
    if ext == ".docx":
        import docx  # python-docx
        d = docx.Document(str(f))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text.strip() for c in r.cells))
        return "\n".join(parts)
    if ext == ".pdf":
        try:
            out = subprocess.run(["pdftotext", "-layout", str(f), "-"], capture_output=True, text=True, timeout=120)
            if out.returncode == 0 and out.stdout.strip():
                return out.stdout
        except (FileNotFoundError, subprocess.TimeoutExpired):
            pass
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(f)).pages)
    return ""


def documents(run, month, files):
    for f in files:
        try:
            text = _extract_text(f)
        except Exception as e:  # keep going; report it
            run.alert(f"documents: could not read {run.rel(f)}: {e}")
            run.recon("documents", month, f, 1, 0, 1)
            run.reject("documents", month, f, 1, f.name, "unreadable")
            continue
        text = run.anon.scrub(text)
        doc_id = "doc_" + short_hash(month, run.rel(f))
        folder = f.parent.name if f.parent.name != month else None
        rec = {"doc_id": doc_id, "period": month, "folder": folder, "file_name": f.name,
               "title": f.stem.replace("_", " "), "doc_type": f.suffix.lower().lstrip("."),
               "initiative": initiative_of(f.name, folder or "", text[:500]), "chars": len(text),
               "text": text, "source_file": run.rel(f), "row_number": 1}
        run.add("documents", rec)
        run.add_study({"study_id": doc_id, "kind": "report", "name": rec["title"], "initiative": rec["initiative"],
                       "period": month, "n": None, "source_file": rec["source_file"]})
        run.recon("documents", month, f, 1, 1, 0)
        run.write_mirror_text(month, f, text)


ADAPTERS = {
    "chargebee": (chargebee_names, chargebee),
    "sprig": (sprig_names, sprig),
    "freshdesk": (None, freshdesk),
    "mbg": (None, mbg),
    "lyssna": (None, lyssna),
    "interviews": (interviews_names, interviews),
}
